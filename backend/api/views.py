from io import BytesIO
from datetime import datetime
from django.core.files.storage import default_storage
from django.db.models import Sum
from django.http import FileResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.urls import reverse
from django_filters.rest_framework import DjangoFilterBackend
from django.db.models.functions import Lower
from docx import Document
from rest_framework import (
    decorators,
    filters,
    permissions,
    status,
    viewsets,
)
from djoser.views import UserViewSet as DjoserUserViewSet

from api.filters import IngredientFilter, RecipeFilter
from api.paginations import Pagination
from api.permissions import IsAuthorOrReadOnly
from api.serializers import (
    AvatarSerializer,
    IngredientSerializer,
    RecipeSerializer,
    UserSerializer,
    UserSubscriptionSerializer,
    RecipeShortSerializer,
)

from recipes.models import (
    Recipe,
    RecipeIngredient,
    Favorite,
    Ingredient,
    User,
    Subscription,
    ShoppingCart
)

from rest_framework.response import Response


def recipe_short_link(request, recipe_id):
    recipe = get_object_or_404(Recipe, id=recipe_id)
    serializer = RecipeShortSerializer(recipe, context={'request': request})
    return JsonResponse(serializer.data)


class IngredientViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = Ingredient.objects.all().order_by(Lower('name'))
    serializer_class = IngredientSerializer
    pagination_class = None
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    filterset_class = IngredientFilter


class RecipeViewSet(viewsets.ModelViewSet):
    queryset = Recipe.objects.all().order_by('-id')
    serializer_class = RecipeSerializer
    permission_classes = [
        permissions.IsAuthenticatedOrReadOnly,
        IsAuthorOrReadOnly,
    ]
    pagination_class = Pagination
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    filterset_class = RecipeFilter

    @decorators.action(detail=True, methods=["get"], url_path="get-link")
    def link(self, request, pk=None):
        if not Recipe.objects.filter(id=pk).exists():
            return Response(
                {"errors": "Рецепт не найден"}, 
                status=status.HTTP_404_NOT_FOUND
            )

        short_url = request.build_absolute_uri(
            reverse('recipe-short-link', kwargs={'recipe_id': pk})
        )
        return Response({"short-link": short_url}, status=status.HTTP_200_OK)

    @decorators.action(
        detail=True,
        methods=["post"],
        url_name="favorite",
        permission_classes=[permissions.IsAuthenticated],
    )
    def favorite(self, request, pk=None):
        return self._add_to_relation(request, pk, Favorite, RecipeShortSerializer)

    @decorators.action(
        detail=True,
        methods=["post"],
        url_name="shopping_cart",
        permission_classes=[permissions.IsAuthenticated],
    )
    def shopping_cart(self, request, pk=None):
        return self._add_to_relation(request, pk, ShoppingCart, RecipeShortSerializer)

    def _add_to_relation(self, request, pk, model, serializer_class):
        recipe = get_object_or_404(Recipe, id=pk)
        user = request.user

        instance, created = model.objects.get_or_create(user=user, recipe=recipe)

        if not created:

            model_name = model._meta.verbose_name
            return Response(
                {"errors": f"Рецепт '{recipe.name}' уже в {model_name}"},
                status=status.HTTP_400_BAD_REQUEST
            )

        serializer = serializer_class(recipe, context={'request': request})
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    def _remove_from_relation(self, request, pk, model):

        user = request.user

        get_object_or_404(model, user=user, recipe_id=pk).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

    @favorite.mapping.delete
    def remove_from_favorite(self, request, pk=None):
        return self._remove_from_relation(request, pk, Favorite)

    @shopping_cart.mapping.delete
    def remove_from_shopping_cart(self, request, pk=None):
        return self._remove_from_relation(request, pk, ShoppingCart)

    @decorators.action(
        detail=False,
        methods=["get"],
        url_path="download_shopping_cart",
        permission_classes=[permissions.IsAuthenticated],
    )
    def download_shopping_cart(self, request):
        document = Document()

        current_date = datetime.now().strftime("%d.%m.%Y %H:%M")
        document.add_heading(f"Список покупок ({current_date})", level=1)

        user_shopping_cart = request.user.recipes_shopping_cart.all()

        document.add_paragraph("Рецепты:")
        for item in user_shopping_cart:
            author_name = item.recipe.author.get_full_name() or item.recipe.author.username
            document.add_paragraph(
                f"- {item.recipe.name} (автор: {author_name})",
                style="ListBullet"
            )

        ingredients = (
            RecipeIngredient.objects.filter(
                recipe__in=user_shopping_cart.values_list("recipe", flat=True)
            )
            .values("ingredient__name", "ingredient__measurement_unit")
            .annotate(total_amount=Sum("amount"))
            .order_by("ingredient__name")
        )

        if ingredients:
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "№"
            hdr_cells[1].text = "Ингредиент"
            hdr_cells[2].text = "Количество"

            for idx, ing in enumerate(ingredients, start=1):
                row_cells = table.add_row().cells
                row_cells[0].text = str(idx)
                name = ing["ingredient__name"].capitalize() if ing["ingredient__name"] else ""
                row_cells[1].text = name
                row_cells[2].text = (
                    f"{ing['total_amount']} "
                    f"{ing['ingredient__measurement_unit']}"
                )
        else:
            document.add_paragraph("Список покупок пуст!")

        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)

        filename = f"shopping_list_{datetime.now().strftime('%Y%m%d')}.docx"

        return FileResponse(
            buffer,
            as_attachment=True,
            filename=filename,
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )


class UserViewSet(DjoserUserViewSet):
    pagination_class = Pagination

    def get_permissions(self):
        if self.action in ['list', 'retrieve']:
            return [permissions.AllowAny()]
        return super().get_permissions()

    def get_serializer_class(self):
        if self.action in ['list', 'retrieve']:
            return UserSerializer
        return super().get_serializer_class()

    @decorators.action(
        detail=False,
        methods=["get"],
        url_path="me",
        permission_classes=[permissions.IsAuthenticated],
    )
    def me(self, request, *args, **kwargs):
        serializer = UserSerializer(
            request.user, context={"request": request}
        )
        return Response(serializer.data)

    @decorators.action(
        detail=False,
        methods=["put", "delete"],
        permission_classes=[permissions.IsAuthenticated],
        url_path="me/avatar",
    )
    def avatar(self, request, *args, **kwargs):
        user = request.user

        if request.method in ["PUT", "PATCH"]:
            serializer = AvatarSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            avatar = serializer.validated_data["avatar"]
            user.avatar = avatar
            user.save()
            return Response(
                {"avatar": user.avatar.url}, status=status.HTTP_200_OK
            )

        if request.method == "DELETE":
            if user.avatar:
                if default_storage.exists(user.avatar.name):
                    default_storage.delete(user.avatar.name)
                user.avatar = None
                user.save()
            return Response(status=status.HTTP_204_NO_CONTENT)

    @decorators.action(
        detail=True,
        methods=["post", "delete"],
        permission_classes=[permissions.IsAuthenticated],
        url_path="subscribe",
    )
    def subscribe(self, request, *args, **kwargs):
        author = self.get_object()
        user = request.user
        if request.method == 'POST':
            if user == author:
                return Response(
                    {'errors': 'Нельзя подписаться на самого себя.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            _, created = Subscription.objects.get_or_create(
                user=user, subscribed_to=author
            )
            if not created:
                return Response(
                    {'errors': f'Вы уже подписаны на пользователя {author.username}.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            data = UserSubscriptionSerializer(
                author, context={'request': request}
            ).data
            return Response(data, status=status.HTTP_201_CREATED)
        get_object_or_404(
            Subscription, user=user, subscribed_to=author
        ).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

    @decorators.action(
        detail=False,
        methods=["get"],
        permission_classes=[permissions.IsAuthenticated],
        url_path="subscriptions",
    )
    def subscriptions(self, request, *args, **kwargs):
        user = request.user
        queryset = User.objects.filter(users_subscribers__user=user)
        pages = self.paginate_queryset(queryset)
        serializer = UserSubscriptionSerializer(
            pages, many=True, context={"request": request}
        )
        return self.get_paginated_response(serializer.data)